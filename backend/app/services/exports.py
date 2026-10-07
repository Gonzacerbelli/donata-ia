import csv
import io
import re
from datetime import UTC, datetime
from typing import Any

from motor.motor_asyncio import AsyncIOMotorDatabase
from openpyxl import Workbook
from openpyxl.styles import Font

from ..core.errors import UnprocessableError
from ..repositories import clients as clients_repo
from ..repositories import products as products_repo
from ..repositories import sales as sales_repo

MAX_EXPORT_ROWS = 10_000

CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SALES_ALIASES = {"sales", "ordenes", "ventas"}
CLIENTS_ALIASES = {"clients", "clientes"}
PRODUCTS_ALIASES = {"products", "productos"}

MONEY_FORMAT = '"$"#,##0'


def _format_date(value: datetime | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).strftime("%d/%m/%Y")


def _payment_status(balance: int, paid: int) -> str:
    if balance <= 0:
        return "pagada"
    if paid <= 0:
        return "sin_pago"
    return "parcial"


def _items_text(items: list) -> str:
    return "; ".join(f"{item.qty} x {item.description}" for item in items)


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


async def _sales_table(
    db: AsyncIOMotorDatabase,
    *,
    status: str | None = None,
    client_id: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    search: str | None = None,
) -> tuple[list[str], list[list[Any]], set[int]]:
    sales = await sales_repo.list_sales(
        db,
        status=status,
        client_id=client_id,
        date_from=date_from,
        date_to=date_to,
        search=search,
    )
    clients = {str(c.id): c for c in await clients_repo.list_clients(db)}
    headers = [
        "Número de orden",
        "Fecha",
        "Cliente",
        "Tipo de cliente",
        "Estado",
        "Estado de cobro",
        "Ítems",
        "Cant. de ítems",
        "Subtotal",
        "Descuento",
        "Costo de envío",
        "Total",
        "Pagado",
        "Saldo",
        "Notas",
    ]
    rows: list[list[Any]] = []
    for sale in sales:
        client = clients.get(str(sale.client_id))
        rows.append(
            [
                sale.id,
                _format_date(sale.date),
                client.name if client else "",
                sale.client_type,
                sale.status,
                _payment_status(sale.balance, sale.paid),
                _items_text(sale.items),
                len(sale.items),
                sale.subtotal,
                sale.discount,
                sale.shipping_cost,
                sale.total,
                sale.paid,
                sale.balance,
                sale.notes or "",
            ]
        )
    return headers, rows, {8, 9, 10, 11, 12, 13}


async def _clients_table(
    db: AsyncIOMotorDatabase,
    *,
    search: str | None = None,
    client_type: str | None = None,
) -> tuple[list[str], list[list[Any]], set[int]]:
    clients = await clients_repo.list_clients(db, search=search, client_type=client_type)
    sales = await sales_repo.list_sales(db)
    totals: dict[str, dict[str, int]] = {}
    for sale in sales:
        if sale.status == "cancelado":
            continue
        acc = totals.setdefault(str(sale.client_id), {"orders": 0, "billed": 0, "balance": 0})
        acc["orders"] += 1
        acc["billed"] += sale.total
        acc["balance"] += sale.balance
    headers = [
        "Nombre",
        "Teléfono",
        "Email",
        "Instagram",
        "Dirección",
        "Tipo",
        "Órdenes",
        "Total facturado",
        "Saldo pendiente",
        "Notas",
        "Fecha de alta",
    ]
    rows: list[list[Any]] = []
    for client in clients:
        acc = totals.get(str(client.id), {"orders": 0, "billed": 0, "balance": 0})
        rows.append(
            [
                client.name,
                client.phone or "",
                client.email or "",
                client.instagram or "",
                client.address or "",
                client.type,
                acc["orders"],
                acc["billed"],
                acc["balance"],
                client.notes or "",
                _format_date(client.created_at),
            ]
        )
    return headers, rows, {7, 8}


async def _products_table(
    db: AsyncIOMotorDatabase,
    *,
    search: str | None = None,
    category: str | None = None,
    provider_id: str | None = None,
    active_only: bool = True,
) -> tuple[list[str], list[list[Any]], set[int]]:
    products = await products_repo.list_products(
        db,
        search=search,
        category=category,
        provider_id=provider_id,
        active_only=active_only,
    )
    providers = {str(p.id): p.name for p in await products_repo.list_providers(db)}
    headers = [
        "Nombre",
        "Categoría",
        "Proveedor",
        "Unidad",
        "Precio minorista",
        "Precio mayorista",
        "Costo",
        "Stock",
        "Stock mínimo",
        "Estado",
    ]
    rows: list[list[Any]] = []
    for product in products:
        rows.append(
            [
                product.name,
                product.category or "",
                providers.get(str(product.provider_id), ""),
                product.unit,
                product.price if product.price is not None else "",
                product.price_mayorista if product.price_mayorista is not None else "",
                product.cost if product.cost is not None else "",
                product.stock,
                product.min_stock,
                "activo" if product.active else "inactivo",
            ]
        )
    return headers, rows, {4, 5, 6}


def _select(filters: dict, *keys: str) -> dict:
    return {key: filters[key] for key in keys if filters.get(key) is not None}


def _csv_cell(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _to_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows([[_csv_cell(cell) for cell in row] for row in rows])
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def _to_xlsx(headers: list[str], rows: list[list[Any]], money_cols: set[int]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for col in range(1, len(headers) + 1):
        sheet.cell(row=1, column=col).font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for row_idx in range(2, len(rows) + 2):
        for col_idx in money_cols:
            sheet.cell(row=row_idx, column=col_idx + 1).number_format = MONEY_FORMAT
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _sales_filename(fmt: str, *, date_from, date_to, status, search) -> str:
    parts = ["ordenes"]
    if date_from:
        parts.append(f"{date_from[:10]}_a_{date_to[:10] if date_to else 'hoy'}")
    if status:
        parts.append(_slug(status))
    if search:
        parts.append(_slug(search))
    return "_".join(parts) + f".{fmt}"


def _clients_filename(fmt: str, *, client_type, search) -> str:
    parts = ["clientes"]
    if client_type:
        parts.append(_slug(client_type))
    if search:
        parts.append(_slug(search))
    return "_".join(parts) + f".{fmt}"


def _products_filename(fmt: str, *, category, search, active_only) -> str:
    parts = ["productos"]
    if category:
        parts.append(_slug(category))
    if search:
        parts.append(_slug(search))
    if not active_only:
        parts.append("incluye-inactivos")
    return "_".join(parts) + f".{fmt}"


async def build_export(
    db: AsyncIOMotorDatabase,
    entity: str,
    fmt: str,
    filters: dict,
) -> tuple[bytes, str, str]:
    if fmt not in {"csv", "xlsx"}:
        raise UnprocessableError("Formato no soportado (usá csv o xlsx)")

    key = entity.lower()
    if key in SALES_ALIASES:
        selected = _select(filters, "status", "client_id", "date_from", "date_to", "search")
        headers, rows, money_cols = await _sales_table(db, **selected)
        filename = _sales_filename(
            fmt,
            date_from=selected.get("date_from"),
            date_to=selected.get("date_to"),
            status=selected.get("status"),
            search=selected.get("search"),
        )
    elif key in CLIENTS_ALIASES:
        selected = _select(filters, "search", "client_type")
        headers, rows, money_cols = await _clients_table(db, **selected)
        filename = _clients_filename(
            fmt, client_type=selected.get("client_type"), search=selected.get("search")
        )
    elif key in PRODUCTS_ALIASES:
        selected = _select(filters, "search", "category", "provider_id", "active_only")
        selected.setdefault("active_only", True)
        headers, rows, money_cols = await _products_table(db, **selected)
        filename = _products_filename(
            fmt,
            category=selected.get("category"),
            search=selected.get("search"),
            active_only=selected.get("active_only", True),
        )
    else:
        raise UnprocessableError("Entidad no exportable")

    if len(rows) > MAX_EXPORT_ROWS:
        raise UnprocessableError(
            f"Demasiados registros ({len(rows)}). Acotá el rango o los filtros."
        )

    if fmt == "csv":
        return _to_csv(headers, rows), filename, CSV_MEDIA_TYPE
    return _to_xlsx(headers, rows, money_cols), filename, XLSX_MEDIA_TYPE
