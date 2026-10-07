import csv
import io
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook


async def _setup(auth_client):
    provider = (await auth_client.post("/providers", json={"name": "Proveedor SA"})).json()
    product = (
        await auth_client.post(
            "/products",
            json={
                "name": "Taza",
                "category": "vajilla",
                "price": 1000,
                "price_mayorista": 800,
                "cost": 500,
                "stock": 20,
                "min_stock": 2,
                "provider_id": provider["id"],
            },
        )
    ).json()
    client = (await auth_client.post("/clients", json={"name": "Ana", "type": "minorista"})).json()
    return provider, product, client


def _rows(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


async def test_sales_csv_has_bom_and_filters(auth_client):
    _, product, client = await _setup(auth_client)
    await auth_client.post(
        "/sales",
        json={"client_id": client["id"], "items": [{"product_id": product["id"], "qty": 2}]},
    )
    other = (
        await auth_client.post(
            "/sales",
            json={"client_id": client["id"], "items": [{"product_id": product["id"], "qty": 1}]},
        )
    ).json()
    await auth_client.patch(f"/sales/{other['id']}", json={"status": "cancelado"})

    response = await auth_client.get("/exports/ordenes.csv", params={"status": "pendiente"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    assert "pendiente" in response.headers["content-disposition"]
    assert response.content.startswith(b"\xef\xbb\xbf")

    rows = _rows(response.content)
    assert rows[0][0] == "Número de orden"
    assert rows[0][6] == "Ítems"
    assert len(rows) == 2
    assert rows[1][2] == "Ana"
    assert rows[1][6] == "2 x Taza"
    assert rows[1][8] == "2000"


async def test_sales_xlsx_money_is_numeric(auth_client):
    _, product, client = await _setup(auth_client)
    await auth_client.post(
        "/sales",
        json={"client_id": client["id"], "items": [{"product_id": product["id"], "qty": 3}]},
    )
    response = await auth_client.get("/exports/ordenes.xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    sheet = load_workbook(io.BytesIO(response.content)).active
    total_cell = sheet.cell(row=2, column=12)
    assert total_cell.value == 3000
    assert total_cell.number_format == '"$"#,##0'


async def test_clients_export_derives_totals(auth_client):
    _, product, client = await _setup(auth_client)
    sale = (
        await auth_client.post(
            "/sales",
            json={"client_id": client["id"], "items": [{"product_id": product["id"], "qty": 2}]},
        )
    ).json()
    await auth_client.post(f"/sales/{sale['id']}/payments", json={"amount": 500})

    response = await auth_client.get("/exports/clientes.csv")
    rows = _rows(response.content)
    assert rows[0][0] == "Nombre"
    assert rows[1][0] == "Ana"
    assert rows[1][6] == "1"
    assert rows[1][7] == "2000"
    assert rows[1][8] == "1500"


async def test_products_export_parity_with_listing(auth_client):
    provider, product, _ = await _setup(auth_client)
    await auth_client.post(
        "/products",
        json={
            "name": "Vaso",
            "category": "cristal",
            "price": 700,
            "cost": 300,
            "stock": 5,
            "min_stock": 1,
            "provider_id": provider["id"],
        },
    )
    listing = await auth_client.get("/products", params={"category": "vajilla"})
    assert len(listing.json()) == 1

    response = await auth_client.get("/exports/productos.csv", params={"category": "vajilla"})
    rows = _rows(response.content)
    assert rows[0][2] == "Proveedor"
    assert len(rows) == 2
    assert rows[1][0] == product["name"]
    assert rows[1][2] == "Proveedor SA"


async def test_csv_escapes_formula_injection(auth_client):
    provider = (await auth_client.post("/providers", json={"name": "Prov"})).json()
    await auth_client.post(
        "/products",
        json={
            "name": "=1+1",
            "price": 100,
            "stock": 1,
            "provider_id": provider["id"],
        },
    )
    response = await auth_client.get("/exports/productos.csv")
    rows = _rows(response.content)
    assert rows[1][0] == "'=1+1"


async def test_xlsx_forces_text_cells(auth_client):
    provider = (await auth_client.post("/providers", json={"name": "Prov"})).json()
    await auth_client.post(
        "/products",
        json={
            "name": "=HYPERLINK(1)",
            "price": 100,
            "stock": 1,
            "provider_id": provider["id"],
        },
    )
    response = await auth_client.get("/exports/productos.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    cell = sheet.cell(row=2, column=1)
    assert cell.value == "=HYPERLINK(1)"
    assert cell.data_type == "s"
    assert cell.number_format == "@"


async def test_export_unknown_entity_returns_422(auth_client):
    assert (await auth_client.get("/exports/foo.csv")).status_code == 422


async def test_export_requires_auth(client):
    assert (await client.get("/exports/ordenes.csv")).status_code == 401


async def test_sales_export_reports_envio_fields(auth_client):
    _, product, client = await _setup(auth_client)
    soon = (datetime.now(UTC) + timedelta(hours=10)).isoformat()
    await auth_client.post(
        "/sales",
        json={
            "client_id": client["id"],
            "items": [{"product_id": product["id"], "qty": 1}],
            "ship_by": soon,
        },
    )
    response = await auth_client.get("/exports/ventas.xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.cell(row=1, column=1).value == "Número de orden"
